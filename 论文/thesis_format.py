# -*- coding: utf-8 -*-
"""毕业论文排版工具（按常见本科规范：A4、小四宋体、1.25倍行距、三线表等）。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, Twips


def set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    run.font.name = en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=0):
    pf = p.paragraph_format
    pf.line_spacing = 1.25
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    p.alignment = align
    if first_line_indent:
        pf.first_line_indent = Cm(0.74)  # 约2字符
    else:
        pf.first_line_indent = Cm(0)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, cn="黑体", en="Times New Roman", size_pt=10.5, bold=True)

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run2 = p.add_run()
    set_run_font(run2, cn="黑体", en="Times New Roman", size_pt=10.5, bold=True)
    run2._r.append(instr)
    run3 = p.add_run()
    set_run_font(run3, cn="黑体", en="Times New Roman", size_pt=10.5, bold=True)
    run3._r.append(fld_char_end)


def set_header(section, title: str):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 页眉下划线
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run(title)
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=10.5, bold=False)


def init_document(title_for_header: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)   # 2cm + 0.5cm 装订
    section.right_margin = Cm(2.0)
    section.gutter = Cm(0)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.line_spacing = 1.25

    set_header(section, title_for_header)
    add_page_number(section)
    return doc


def add_center_title(doc, text, size_pt=18, cn="黑体", en="Times New Roman", bold=True, space_before=0, space_after=12):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=space_before, space_after=space_after)
    run = p.add_run(text)
    set_run_font(run, cn=cn, en=en, size_pt=size_pt, bold=bold)
    return p


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=first_indent, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    # 支持简单上标引用：把 [n] 做成上标
    import re
    parts = re.split(r"(\[\d+\])", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run(part)
        set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)
        if re.fullmatch(r"\[\d+\]", part):
            run.font.superscript = True
    return p


def add_heading1(doc, text):
    """一级标题：四号黑体加粗，缩进2字符，单独成行。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6)
    run = p.add_run(text)
    set_run_font(run, cn="黑体", en="Times New Roman", size_pt=14, bold=True)
    return p


def add_heading2(doc, text):
    """二级及以下：小四宋体加粗。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    run = p.add_run(text)
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12, bold=True)
    return p


def add_heading3(doc, text):
    return add_heading2(doc, text)


def add_blank(doc):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    return p


def add_caption(doc, text, above=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6 if not above else 6, space_after=3)
    run = p.add_run(text)
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12, bold=True)
    return p


def _set_cell_border(cell, **edges):
    """edges: top/bottom/left/right -> (sz, val) sz in eighths of a point."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        if edge in edges and edges[edge]:
            sz, val = edges[edge]
            element.set(qn("w:val"), val)
            element.set(qn("w:sz"), str(sz))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")
        else:
            element.set(qn("w:val"), "nil")
        tcBorders.append(element)
    tcPr.append(tcBorders)


def add_three_line_table(doc, headers, rows, col_widths=None):
    """三线表：顶线、表头下线、底线。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    def fill_cell(cell, text, bold=False, size=10.5):
        cell.text = ""
        p = cell.paragraphs[0]
        set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(str(text))
        set_run_font(run, cn="宋体", en="Times New Roman", size_pt=size, bold=bold)

    # header
    for i, h in enumerate(headers):
        fill_cell(table.rows[0].cells[i], h, bold=True)
        _set_cell_border(
            table.rows[0].cells[i],
            top=(12, "single"),
            bottom=(6, "single"),
            left=None,
            right=None,
        )

    # body
    last = len(rows)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            fill_cell(table.rows[r_idx + 1].cells[c_idx], val, bold=False)
            bottom = (12, "single") if r_idx == last - 1 else None
            _set_cell_border(
                table.rows[r_idx + 1].cells[c_idx],
                top=None,
                bottom=bottom,
                left=None,
                right=None,
            )

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Cm(w)
    return table


def add_toc_line(doc, title, level=1):
    """手工目录行（页码请在 Word 中更新或手填）。"""
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    indent = (level - 1) * 0.74
    p.paragraph_format.left_indent = Cm(indent)
    # 点线页码：用制表符
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = p.add_run(title)
    bold = level == 1
    set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12, bold=bold)
    p.add_run("\t")
    run2 = p.add_run("")  # 页码占位，Word里可再填
    set_run_font(run2, cn="宋体", en="Times New Roman", size_pt=12)
    return p


def page_break(doc):
    doc.add_page_break()
