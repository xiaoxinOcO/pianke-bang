# -*- coding: utf-8 -*-
"""
生成《基于SpringCloud+Vue的校园便民跑腿平台的设计与实现》毕业论文 Word 文档。
仅在「论文」目录产出文档，不修改业务代码。

用法：
  python build_thesis.py
"""

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from thesis_content import (
    CN_ABSTRACT,
    CN_KEYWORDS,
    EN_ABSTRACT,
    EN_KEYWORDS,
    META,
    REFERENCES,
    TOC_ENTRIES,
)
from thesis_format import (
    add_blank,
    add_body,
    add_caption,
    add_center_title,
    add_heading1,
    add_heading2,
    add_heading3,
    add_three_line_table,
    add_toc_line,
    init_document,
    page_break,
    set_paragraph_format,
    set_run_font,
)

OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / f"{META['title']}.docx"


def add_cover(doc):
    for _ in range(2):
        add_blank(doc)
    add_center_title(doc, META["school"], size_pt=22, space_after=6)
    add_center_title(doc, "本科毕业论文（设计）", size_pt=18, space_after=24)
    add_blank(doc)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    run = p.add_run(META["title"])
    set_run_font(run, cn="黑体", en="Times New Roman", size_pt=16, bold=True)

    add_blank(doc)
    add_blank(doc)

    rows = [
        ("学　　院", META["college"]),
        ("专　　业", META["major"]),
        ("学生姓名", META["student"]),
        ("学　　号", META["student_id"]),
        ("指导教师", f"{META['advisor']}（{META['advisor_title']}）"),
        ("完成时间", f"{META['year']}年"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        run = p.add_run(f"{label}：{value}")
        set_run_font(run, cn="宋体", en="Times New Roman", size_pt=14, bold=False)
    page_break(doc)


def add_cn_abstract(doc):
    add_center_title(doc, "摘　　要", size_pt=18, space_before=12, space_after=12)
    add_blank(doc)
    add_body(doc, CN_ABSTRACT)
    add_blank(doc)
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    r1 = p.add_run("关键词：")
    set_run_font(r1, cn="宋体", en="Times New Roman", size_pt=14, bold=True)
    r2 = p.add_run(CN_KEYWORDS)
    set_run_font(r2, cn="宋体", en="Times New Roman", size_pt=12, bold=False)
    page_break(doc)


def add_en_abstract(doc):
    add_center_title(doc, "Abstract", size_pt=18, cn="Times New Roman", en="Times New Roman", space_before=12, space_after=12)
    add_blank(doc)
    add_body(doc, EN_ABSTRACT)
    add_blank(doc)
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    r1 = p.add_run("Key words: ")
    set_run_font(r1, cn="Times New Roman", en="Times New Roman", size_pt=14, bold=True)
    r2 = p.add_run(EN_KEYWORDS)
    set_run_font(r2, cn="Times New Roman", en="Times New Roman", size_pt=12, bold=False)
    page_break(doc)


def add_toc(doc):
    add_center_title(doc, "目　　录", size_pt=18, space_before=12, space_after=12)
    for title, level in TOC_ENTRIES:
        add_toc_line(doc, title, level=level)
    page_break(doc)


def add_chapter1(doc):
    add_heading1(doc, "1 绪论")
    add_heading2(doc, "1.1 研究背景与意义")
    add_body(
        doc,
        "随着高校校园规模不断扩大，学生学习与生活节奏明显加快，快递代取、物品代送、资料代拿、代买餐食等互助型跑腿需求持续增长。"
        "智慧校园建设正从单一业务系统走向数据共享、业务互通与一站式服务，学生对便捷、可信、低成本的校园互助服务需求更加突出[1]。"
        "与此同时，移动互联网推动校园生活服务向平台化、数字化转型，传统微信群接龙、口头约定等方式，已难以满足任务发布、接单匹配、"
        "资金结算与纠纷处理的规范化需要[2]。校园快递与末端配送场景中仍普遍存在堆积、取件时间冲突、人工协调成本高、履约过程难追溯等问题，"
        "影响用户体验与平台精细化运营[3]。因此，研究并实现一套基于SpringBoot与SpringCloud的校园便民跑腿平台，具有现实应用价值。",
    )
    add_body(
        doc,
        "从工程实践角度看，基于SpringCloud微服务架构的系统开发，通过服务拆分与Nacos等组件实现低耦合、可扩展部署，"
        "为校园跑腿平台在认证、任务、钱包、聊天、审核等模块的独立协同运行提供了技术路径[4]。"
        "前后端分离模式下，后端以SpringBoot提供接口、前端以Vue构建多端界面，可提升开发效率与可维护性[5]。"
        "本课题以“片刻帮”为产品形态，将任务撮合、资金结算、信誉约束、实时沟通与客服申诉流程化，"
        "既服务师生日常互助，也为软件工程专业综合运用微服务、前后端分离与分布式中间件提供完整案例。",
    )

    add_heading2(doc, "1.2 国内外研究现状")
    add_body(
        doc,
        "国外对众包配送与最后一公里协同配送研究起步较早。相关工作讨论了众包配送在规划与运营层面的挑战与机会，"
        "强调运力波动、匹配效率与服务质量控制等关键问题[6]。面向大学场景的零工经济配送模式研究指出，"
        "校园相对封闭、用户同质性强，适合开展低成本、短距离的互助配送试点[7]。"
        "另有研究评估公共交通结合众包配送的潜在影响，为校园与城市末端协同提供了参考[8]。",
    )
    add_body(
        doc,
        "国内在智慧校园、校园跑腿小程序、快递末端及微服务业务系统方面成果较多。"
        "有研究从校园跑腿小程序角度探讨发布、接单与评价流程，验证了移动端互助服务的可用性[2]；"
        "亦有工作面向校园快递智能无人车与配送机器人开展系统设计与路径规划[3][10]；"
        "众包物流视角下的校园快递末端配送模式研究强调运力组织与交付方式对效率的影响[9][11]。"
        "在工程实现方面，微服务与SpringBoot+Vue在高校信息系统中已有较多落地经验[4][5]，"
        "智慧校园后勤信息化平台构建亦强调业务中台化与服务整合[12]。"
        "综合来看，已有研究多聚焦某一侧能力，一体化覆盖多类型跑腿、钱包结算、信誉申诉、实时通信、智能客服与审核的方案仍相对不足。"
        "本文据此构建多角色微服务平台，补齐任务、资金、沟通与运营治理闭环。",
    )

    add_heading2(doc, "1.3 研究内容与方法")
    add_body(
        doc,
        "本文主要研究内容包括：（1）围绕校园跑腿场景开展可行性与需求分析，明确管理员、普通用户、接单员与客服四类角色职责；"
        "（2）基于SpringBoot与SpringCloud完成系统总体架构、功能模块、数据库与接口详细设计；"
        "（3）实现网关、认证、用户、任务、审核、聊天、钱包等微服务，以及用户主站、接单员工作台、客服工作台三端前端；"
        "（4）对登录鉴权、任务主流程、支付提现、即时通信、管理审核等进行功能、性能与安全性测试。",
    )
    add_body(
        doc,
        "研究方法采用文献研究法梳理国内外相关成果；采用问卷调查与场景访谈提炼功能与非功能需求；"
        "采用原型法与迭代开发完成系统实现；采用黑盒测试与接口联调验证系统正确性与稳定性。",
    )

    add_heading2(doc, "1.4 论文组织结构")
    add_body(
        doc,
        "第1章介绍研究背景、国内外现状、研究内容与方法；第2章介绍SpringBoot、SpringCloud Alibaba、Vue、MySQL及关键中间件；"
        "第3章给出可行性、功能与非功能需求分析；第4章给出总体架构、数据库与详细设计；第5章说明各服务与前端实现；"
        "第6章给出测试环境、用例与结论；第7章总结全文并展望后续工作。",
    )
    page_break(doc)


def add_chapter2(doc):
    add_heading1(doc, "2 相关技术介绍")
    add_heading2(doc, "2.1 SpringBoot")
    add_body(
        doc,
        "SpringBoot通过约定优于配置、内嵌容器与起步依赖，显著降低企业级Java应用搭建成本。"
        "本系统各微服务基于SpringBoot 3.2.x构建，统一提供REST接口、配置管理与健康检查能力，便于在微服务环境中快速迭代。",
    )
    add_heading2(doc, "2.2 SpringCloud Alibaba")
    add_body(
        doc,
        "SpringCloud Alibaba提供面向微服务的一站式组件体系。本系统采用2023.x相关版本，结合Nacos、Gateway、Sentinel、Seata等能力，"
        "支撑服务注册发现、统一入口、流量防护与分布式事务，契合校园跑腿平台多业务模块协同的需求[4]。",
    )
    add_heading2(doc, "2.3 Vue.js")
    add_body(
        doc,
        "Vue.js是渐进式前端框架，组件化与响应式数据绑定适合构建中大型单页应用。"
        "本系统用户主站、接单员工作台与客服工作台均基于Vue 3与Vite实现，通过统一API网关访问后端服务，实现前后端分离[5]。",
    )
    add_heading2(doc, "2.4 MySQL")
    add_body(
        doc,
        "MySQL是成熟的关系型数据库。本系统使用MySQL 8.x存储用户、任务、钱包、聊天、优惠券、公告等核心业务数据，"
        "库名为campus_delivery，并通过JPA与迁移脚本管理表结构演进。",
    )
    add_heading2(doc, "2.5 Nacos与SpringCloud Gateway")
    add_body(
        doc,
        "Nacos承担服务注册发现与配置中心职责，使各微服务可动态上下线。"
        "SpringCloud Gateway作为统一入口（默认8080端口），对/api/**与/ws/**进行路由转发，并完成JWT鉴权与跨域处理，降低前端对接复杂度。",
    )
    add_heading2(doc, "2.6 OpenFeign、Sentinel与Seata")
    add_body(
        doc,
        "OpenFeign用于任务服务调用钱包服务等跨服务通信；Sentinel用于网关层流控与熔断，提升高峰可用性；"
        "Seata用于赏金锁定、退款等跨服务事务场景，保障资金与任务状态一致。",
    )
    add_heading2(doc, "2.7 RabbitMQ与WebSocket")
    add_body(
        doc,
        "RabbitMQ用于任务状态变更等异步通知，降低同步链路耦合；WebSocket用于任务私聊与客服会话的实时推送，"
        "满足校园场景下即时沟通需求。",
    )
    add_heading2(doc, "2.8 支付宝开放平台与Ollama")
    add_body(
        doc,
        "系统对接支付宝沙箱完成充值、赏金支付与提现相关流程，便于演示与验收。"
        "Ollama用于本地或内网大模型推理，支撑任务内容智能审核配置与智能客服问答，减轻人工审核与一线客服压力。",
    )
    page_break(doc)


def add_chapter3(doc):
    add_heading1(doc, "3 系统分析")
    add_heading2(doc, "3.1 可行性分析")
    add_heading3(doc, "3.1.1 经济可行性")
    add_body(
        doc,
        "校园跑腿需求稳定、用户集中，平台可采用轻资产运营。开发阶段主要依赖开源框架与本地/容器化中间件，"
        "硬件与第三方服务成本可控，具备经济可行性。",
    )
    add_heading3(doc, "3.1.2 技术可行性")
    add_body(
        doc,
        "JDK 17、SpringBoot、SpringCloud Alibaba、Vue 3、MySQL及Nacos、RabbitMQ、Seata等技术栈成熟，"
        "团队已具备本地联调与分模块开发经验，技术风险可控。",
    )
    add_heading3(doc, "3.1.3 操作可行性")
    add_body(
        doc,
        "系统按角色提供用户主站、接单员工作台、客服工作台与管理后台，界面路径清晰，"
        "普通用户可快速完成注册、发布与接单，管理员可在后台完成审核与运营，操作可行性良好。",
    )

    add_heading2(doc, "3.2 功能需求分析")
    add_heading3(doc, "3.2.1 认证与用户中心模块")
    add_body(
        doc,
        "支持密码登录与短信验证码注册/登录/找回密码；支持个人资料与头像维护；支持接单员申请与认证状态查看；"
        "支持信誉分与评价信息展示，为后续接单与奖惩提供依据。",
    )
    add_heading3(doc, "3.2.2 任务发布与接单模块")
    add_body(
        doc,
        "发布者可创建代取、代买等类型任务并支付赏金；任务经审核后进入大厅；接单员可浏览、接单、提交完成凭证；"
        "发布者可确认、评价；支持取消、申诉裁决、加急与统计看板等运营能力。",
    )
    add_heading3(doc, "3.2.3 钱包支付与优惠券模块")
    add_body(
        doc,
        "提供余额账户、流水查询、支付宝充值、赏金支付、提现申请与审核；支持优惠券领取/兑换与大厅活动，"
        "实现营销触达与支付成本降低。",
    )
    add_heading3(doc, "3.2.4 实时聊天与消息模块")
    add_body(
        doc,
        "任务双方可私聊沟通取件细节；支持已读回执；系统消息与公告用于通知审核结果、资金变动与平台规则；"
        "用户可提交反馈。",
    )
    add_heading3(doc, "3.2.5 智能客服与人工客服模块")
    add_body(
        doc,
        "用户可通过智能客服咨询钱包、任务等高频问题；必要时转人工；客服工作台提供会话列表、快捷回复与订单上下文；"
        "支持客服评价，形成服务闭环。",
    )
    add_heading3(doc, "3.2.6 管理员模块")
    add_body(
        doc,
        "管理员可管理用户启停、头像审核、接单员审核、信誉分；审核任务与处理申诉；审核提现与开展资金运营；"
        "管理公告、日志、反馈、客服账号与AI配置等。",
    )

    add_heading2(doc, "3.3 非功能需求分析")
    add_heading3(doc, "3.3.1 性能需求")
    add_body(
        doc,
        "常规业务接口应在可接受时延内返回；任务大厅、登录等热点路径应支持基础并发；"
        "即时消息应具备较低延迟；可通过限流与异步消息削峰。",
    )
    add_heading3(doc, "3.3.2 安全性需求")
    add_body(
        doc,
        "需实现JWT统一鉴权与角色授权；密码不可明文存储；管理接口严格鉴权；支付与提现过程可审计；"
        "任务内容可审核，降低违规信息风险。",
    )
    add_heading3(doc, "3.3.3 可维护性与可扩展性需求")
    add_body(
        doc,
        "微服务应可独立部署；前后端分离便于并行开发；配置集中于Nacos；数据库表结构可通过迁移演进，"
        "便于后续扩展小程序端或新业务类型。",
    )
    page_break(doc)


def add_chapter4(doc):
    add_heading1(doc, "4 系统设计")
    add_heading2(doc, "4.1 系统总体设计")
    add_heading3(doc, "4.1.1 系统架构设计")
    add_body(
        doc,
        "系统采用“前端多端 + API网关 + 业务微服务 + 中间件 + MySQL”分层架构。"
        "表现层包括frontend用户主站与管理壳、rider-console接单员工作台、cs-console客服工作台；"
        "网关层由campus-gateway负责路由、JWT鉴权、跨域与Sentinel防护；"
        "服务层包括auth、user、task、audit、chat、wallet等微服务；"
        "支撑层包括Nacos、RabbitMQ、Seata及可选Elasticsearch任务统计；数据层为MySQL数据库campus_delivery。"
        "总体架构示意如图4-1所示（可在定稿时插入架构截图或Visio图）。",
    )
    add_caption(doc, "图4-1 系统总体架构图（请插入）")

    add_heading3(doc, "4.1.2 功能模块划分")
    add_body(
        doc,
        "按微服务职责划分：认证服务负责注册登录与令牌签发；用户服务负责用户治理与接单员/客服账号管理；"
        "任务服务负责任务全生命周期；钱包服务负责余额、支付、提现、优惠券与活动；聊天服务负责即时消息；"
        "审核服务负责公告、反馈、智能客服与AI配置等。公共能力沉淀于campus-common与campus-persistence。",
    )
    add_caption(doc, "表4-1 微服务模块划分")
    add_three_line_table(
        doc,
        ["模块", "端口/入口", "主要职责"],
        [
            ["campus-gateway", "8080", "统一路由、JWT鉴权、限流"],
            ["campus-auth-service", "8101+", "注册登录、短信、资料、接单申请"],
            ["campus-user-service", "8101+", "用户管理、信誉、客服账号"],
            ["campus-task-service", "8101+", "任务发布/接单/确认/申诉"],
            ["campus-wallet-service", "8101+", "钱包、支付宝、优惠券、活动"],
            ["campus-chat-service", "8101+", "私聊、WebSocket推送"],
            ["campus-audit-service", "8101+", "公告、客服、AI配置、日志"],
        ],
        col_widths=[4.5, 3.0, 7.0],
    )

    add_heading2(doc, "4.2 数据库设计")
    add_heading3(doc, "4.2.1 E-R图")
    add_body(
        doc,
        "核心实体包括用户、任务、钱包账户、钱包流水、提现申请、聊天消息、系统消息、公告、反馈、"
        "优惠券模板/用户券、大厅活动、客服评价等。用户与任务为一对多；用户与钱包账户为一对一；"
        "任务与流水、聊天记录关联。E-R图如图4-2所示（定稿插入）。",
    )
    add_caption(doc, "图4-2 系统E-R图（请插入）")

    add_heading3(doc, "4.2.2 数据表设计")
    add_body(
        doc,
        "结合系统实现，主要数据表如表4-2所示。论文定稿时可再展开用户表、任务表、钱包账户表等字段级三线表。",
    )
    add_caption(doc, "表4-2 主要数据表一览")
    add_three_line_table(
        doc,
        ["表名", "说明"],
        [
            ["user（用户表）", "账号、密码摘要、手机号、角色/接单认证、信誉分、头像等"],
            ["task（任务表）", "类型、标题、赏金、状态、发布者/接单者、加急、凭证等"],
            ["wallet_account", "用户余额账户"],
            ["wallet_transaction", "收支流水"],
            ["wallet_withdraw_request", "提现申请与审核状态"],
            ["chat相关表", "消息内容、会话双方、已读游标等"],
            ["user_system_message", "系统通知"],
            ["user_feedback", "用户反馈"],
            ["coupon_template / user_coupon", "优惠券模板与用户券"],
            ["hall_activity", "大厅活动"],
            ["cs_rating", "客服评价"],
            ["sms_verification_code", "短信验证码"],
        ],
        col_widths=[5.5, 9.0],
    )

    add_caption(doc, "表4-3 任务表（task）主要字段")
    add_three_line_table(
        doc,
        ["字段", "类型", "说明"],
        [
            ["id", "BIGINT", "主键"],
            ["title", "VARCHAR", "任务标题"],
            ["type", "VARCHAR", "任务类型"],
            ["bounty", "DECIMAL", "赏金金额"],
            ["status", "VARCHAR", "任务状态"],
            ["publisher_id", "BIGINT", "发布者用户ID"],
            ["rider_id", "BIGINT", "接单员用户ID"],
            ["urgent", "TINYINT", "是否加急"],
            ["created_at", "DATETIME", "创建时间"],
        ],
        col_widths=[4.0, 3.5, 7.0],
    )

    add_heading2(doc, "4.3 详细设计")
    add_heading3(doc, "4.3.1 前端界面设计")
    add_body(
        doc,
        "用户主站包含登录注册、任务大厅、任务发布、我的任务、钱包与优惠券、聊天、个人中心、智能客服及管理后台入口；"
        "接单员工作台突出任务大厅、我的订单、钱包提现与任务私聊；客服工作台突出会话列表、快捷回复与用户/订单上下文面板。"
        "三端统一通过网关访问后端，保持交互风格一致。",
    )
    add_heading3(doc, "4.3.2 后端接口设计")
    add_body(doc, "后端接口按业务域分组，主要路径前缀如表4-4所示。")
    add_caption(doc, "表4-4 后端接口分组")
    add_three_line_table(
        doc,
        ["分组", "路径前缀", "说明"],
        [
            ["认证", "/api/auth/**", "注册登录、短信、资料、接单申请"],
            ["任务", "/api/tasks/**", "发布/接单/完成/确认/评价/申诉"],
            ["钱包", "/api/wallet/**", "余额、提现、支付宝相关"],
            ["优惠券/活动", "/api/coupons/** 等", "领券与大厅活动"],
            ["聊天", "/api/chat/**、/ws/**", "消息与WebSocket"],
            ["客服", "/api/support/**", "智能客服、转人工、评价"],
            ["管理", "/api/admin/**", "用户/任务/资金/公告/AI等"],
        ],
        col_widths=[3.5, 5.0, 6.0],
    )

    add_heading3(doc, "4.3.3 网关鉴权与分布式事务设计")
    add_body(
        doc,
        "客户端请求首先到达Gateway，校验JWT有效性并将用户身份传递给下游服务；未携带令牌或令牌失效的受保护接口将被拒绝。"
        "任务发布并支付赏金时，任务服务通过OpenFeign调用钱包服务锁定余额，并以Seata保障跨服务一致性；"
        "失败则回滚/退款。任务状态变更可通过RabbitMQ异步通知相关方，降低同步链路压力。",
    )
    page_break(doc)


def add_chapter5(doc):
    add_heading1(doc, "5 系统实现")
    add_heading2(doc, "5.1 开发环境搭建")
    add_body(
        doc,
        "开发环境为Windows 11，JDK 17，Maven 3.9+，IntelliJ IDEA，Node.js与Vite，MySQL 8。"
        "通过Docker Compose部署Nacos、RabbitMQ、Seata等中间件。"
        "后端父工程campus-cloud聚合各微服务；前端工程包括frontend、rider-console、cs-console。"
        "本地联调时，前端开发服务器将/api代理至网关8080端口。",
    )

    add_heading2(doc, "5.2 网关与认证服务实现")
    add_body(
        doc,
        "campus-gateway按Path将请求路由到认证、用户、任务、钱包、聊天、审核等服务，并对/ws/**提供WebSocket转发。"
        "campus-auth-service实现/api/auth下注册、登录、短信验证码、忘记密码、个人资料与接单员申请等接口，"
        "签发JWT供网关与下游识别身份。短信能力可对接阿里云号码认证，开发环境支持模拟验证码以便联调。",
    )

    add_heading2(doc, "5.3 用户服务实现")
    add_body(
        doc,
        "campus-user-service面向管理端提供用户列表、启停、头像审核、接单员通过/驳回/撤销、信誉分调整、"
        "客服账号管理及客服分配相关能力，支撑平台治理与风控。",
    )

    add_heading2(doc, "5.4 任务服务实现")
    add_body(
        doc,
        "campus-task-service实现任务发布、大厅查询、接单、完成、确认、评价、取消、审核通过/驳回、重提、申诉裁决等接口，"
        "并与钱包内部接口协同完成赏金锁定与结算。支持任务统计看板与AI批量审核入口，"
        "可按配置在赏金支付成功后进入待审或自动过审策略。",
    )

    add_heading2(doc, "5.5 钱包服务实现")
    add_body(
        doc,
        "campus-wallet-service实现余额查询、提现申请与审核、支付宝充值/支付相关接口，"
        "以及优惠券模板发放、兑换码与大厅活动管理。对任务服务暴露lock-bounty、refund-bounty、"
        "income-from-task等内部接口，保证赏金流转可追踪。",
    )

    add_heading2(doc, "5.6 聊天服务实现")
    add_body(
        doc,
        "campus-chat-service提供发消息、历史记录、联系人、已读回执等REST接口，"
        "并通过WebSocket向在线用户推送实时消息，支撑任务沟通与客服会话。",
    )

    add_heading2(doc, "5.7 审核与客服服务实现")
    add_body(
        doc,
        "campus-audit-service承载公告、系统消息、反馈、智能客服/api/support/chat、客服评价、操作日志、"
        "AI配置测试等能力。智能客服结合知识检索与只读工具查询钱包/任务信息，必要时引导转人工，"
        "由cs-console完成后续处理。",
    )

    add_heading2(doc, "5.8 前端三端实现")
    add_body(
        doc,
        "用户主站实现大厅、发布、钱包、聊天、客服与管理后台等页面；接单员工作台实现接单履约与提现；"
        "客服工作台实现会话处理与快捷回复。三端统一经网关访问后端API，并根据角色控制菜单与路由。"
        "关键界面截图可置于本节（钱包支付、支付宝沙箱等素材已在论文目录备选）。",
    )
    page_break(doc)


def add_chapter6(doc):
    add_heading1(doc, "6 系统测试")
    add_heading2(doc, "6.1 测试环境与方法")
    add_body(
        doc,
        "测试环境与开发环境基本一致。采用黑盒功能测试、接口联调、基础并发访问与权限用例验证。"
        "重点覆盖登录鉴权、任务主流程、支付提现、聊天推送与管理审核。",
    )

    add_heading2(doc, "6.2 功能测试")
    add_body(doc, "主要功能测试用例如表6-1所示。")
    add_caption(doc, "表6-1 功能测试用例")
    add_three_line_table(
        doc,
        ["编号", "用例", "预期结果", "结果"],
        [
            ["F1", "短信/密码登录", "返回令牌并进入对应端", "通过"],
            ["F2", "发布任务并支付赏金", "任务进入待审/大厅，余额正确", "通过"],
            ["F3", "接单并提交完成", "状态流转正确，可确认评价", "通过"],
            ["F4", "提现申请与审核", "流水与余额一致", "通过"],
            ["F5", "WebSocket私聊", "双方实时收到消息", "通过"],
            ["F6", "智能客服转人工", "客服台可见并可回复", "通过"],
        ],
        col_widths=[2.0, 4.0, 5.5, 2.0],
    )

    add_heading2(doc, "6.3 性能测试")
    add_body(
        doc,
        "对大厅列表、登录、接单等接口进行基础并发访问验证。在超限情况下，Gateway结合Sentinel可返回限流提示，"
        "系统整体保持可用，满足课程设计与毕设验收层面的性能要求。更系统的压测可作为后续工作。",
    )

    add_heading2(doc, "6.4 安全性测试")
    add_body(
        doc,
        "无令牌访问受保护接口应被拒绝；普通用户无法调用管理接口；未通过接单认证的用户不可接单；"
        "密码不以明文存储；越权操作失败。测试结果表明鉴权与授权机制有效。",
    )

    add_heading2(doc, "6.5 测试结论")
    add_body(
        doc,
        "主要功能符合需求分析，关键资金与任务状态流转正确，鉴权与基本性能满足验收要求。"
        "后续可加强自动化回归、长稳压测与安全加固。",
    )
    page_break(doc)


def add_chapter7(doc):
    add_heading1(doc, "7 总结与展望")
    add_heading2(doc, "7.1 总结")
    add_body(
        doc,
        "本文完成了基于SpringBoot与SpringCloud的校园便民跑腿平台设计与实现，形成网关+多微服务+三端前端的完整方案，"
        "覆盖任务、钱包、聊天、客服与运营管理等核心业务，验证了微服务架构在校园互助场景的可行性，"
        "也为软件工程综合实践提供了可复现的工程案例。",
    )
    add_heading2(doc, "7.2 展望")
    add_body(
        doc,
        "后续可完善移动端小程序体验，探索智能派单与路径推荐，完善监控告警与更细粒度权限，"
        "并扩大真实校园试运行与安全加固，持续提升系统可用性与用户体验。",
    )
    page_break(doc)


def add_references(doc):
    add_center_title(doc, "参考文献", size_pt=18, space_before=12, space_after=12)
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        run = p.add_run(f"[{i}] {ref}")
        set_run_font(run, cn="宋体", en="Times New Roman", size_pt=12)
    page_break(doc)


def add_thanks(doc):
    add_center_title(doc, "致　　谢", size_pt=18, space_before=12, space_after=12)
    add_blank(doc)
    add_body(
        doc,
        f"本论文是在指导教师{META['advisor']}老师的悉心指导下完成的。老师在选题论证、架构设计、论文撰写与修改完善等方面给予了耐心指导与宝贵建议，"
        "使我能够顺利完成本课题的研究与系统实现。在此谨向老师致以诚挚的感谢。",
    )
    add_body(
        doc,
        "感谢广州工商学院AI工程师学院各位老师在本科阶段的培养与帮助，感谢同学与朋友在联调测试、资料整理过程中提供的支持。"
        "由于本人水平有限，论文与系统仍有不足之处，恳请各位老师批评指正。",
    )


def build():
    doc = init_document(META["title"])
    add_cover(doc)
    add_cn_abstract(doc)
    add_en_abstract(doc)
    add_toc(doc)
    add_chapter1(doc)
    add_chapter2(doc)
    add_chapter3(doc)
    add_chapter4(doc)
    add_chapter5(doc)
    add_chapter6(doc)
    add_chapter7(doc)
    add_references(doc)
    add_thanks(doc)
    doc.save(str(OUT_FILE))
    print(f"OK: {OUT_FILE}")


if __name__ == "__main__":
    build()
